# -*- coding: utf-8 -*-
"""在**模板副本**上填内容：严格沿用模板表格设计，只填文本 / 换图片。

重点：用例侧重 UI 界面与实跑检测；数量不设 50 上限，超出部分按模板样式自动扩展行。

来源：第xx小组_xxx项目_测试成绩书.xlsx
输出：docs/testbook/第xx小组_垃圾自动识别分类系统_测试成绩书.xlsx
"""
from __future__ import annotations

import shutil
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

from cases_final import CASES_FINAL

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
TEMPLATE = ROOT / "第xx小组_xxx项目_测试成绩书.xlsx"
OUT = HERE / "第xx小组_垃圾自动识别分类系统_测试成绩书.xlsx"
SHOTS = HERE / "shots"

PROJECT = "垃圾自动识别分类系统"
TODAY = date.today()
FIRST_ROW = 5        # 模板数据行起点
TEMPLATE_LAST_ROW = 54   # 模板原有数据行终点（50 行）

# 模板的数据列分组（每组是一个合并区域）
GROUPS = [("B", "D"), ("E", "U"), ("V", "AL"), ("AM", "AP"),
          ("AQ", "AT"), ("AU", "AX"), ("AY", "BB"), ("BC", "BF")]
COL_RANGE = [get_column_letter(c) for c in range(2, 59)]  # B..BF


def fill_cover(wb) -> None:
    ws = wb["文件封面"]
    ws["A3"] = "文件编号：TS-2026-0001"
    ws["A15"] = f"{PROJECT}测试成绩书"
    ws["A17"] = "版本：1.0.0"
    ws["A18"] = TODAY


def ensure_data_rows(ws, count: int) -> int:
    """把数据行补足到 count 行：复制模板最后一行的样式，并补上同样的合并区域。"""
    need = FIRST_ROW + count - 1
    if need <= TEMPLATE_LAST_ROW:
        return 0
    added = 0
    src_row = TEMPLATE_LAST_ROW
    for r in range(TEMPLATE_LAST_ROW + 1, need + 1):
        for col in COL_RANGE:
            src = ws[f"{col}{src_row}"]
            dst = ws[f"{col}{r}"]
            dst._style = copy(src._style)
            dst.border = copy(src.border)
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.alignment = copy(src.alignment)
            dst.protection = copy(src.protection)
            dst.number_format = src.number_format
        ws.row_dimensions[r].height = ws.row_dimensions[src_row].height
        for a, b in GROUPS:
            ws.merge_cells(f"{a}{r}:{b}{r}")
        added += 1
    return added


def evidence_of(case: tuple) -> str:
    ev = (case[6] or "").strip()
    if ev.startswith("见"):
        return ev
    mod = case[0]
    if mod == "UI界面":
        return "见UI界面全景（标注编号）"
    if mod == "检测运行":
        return "见运行图"
    return ev or "见运行记录"


def fill_cases(wb) -> int:
    ws = wb["测试项目"]
    ws["B1"] = f"{PROJECT}项目单体测试书"
    ws["X2"] = PROJECT
    ws["AQ3"] = TODAY.strftime("%Y/%m")

    rows = CASES_FINAL
    ensure_data_rows(ws, len(rows))

    for i, case in enumerate(rows):
        r = FIRST_ROW + i
        mod, title, pre, steps, exp, method, _ev, _res, verdict, _note = case
        ws[f"B{r}"] = i + 1
        ws[f"E{r}"] = f"用例标题：{title}\n前置条件：{pre}\n操作步骤：{steps}"
        ws[f"V{r}"] = exp
        ws[f"AM{r}"] = verdict
        ws[f"AQ{r}"] = TODAY
        ws[f"AU{r}"] = None
        ws[f"AY{r}"] = None
        ws[f"BC{r}"] = evidence_of(case)
    return len(rows)


def add_img(ws, path: Path, cell: str, width: int) -> None:
    if not path.exists():
        print(f"  [跳过] 缺少图片 {path.name}")
        return
    img = XLImage(str(path))
    scale = width / img.width
    img.width, img.height = width, int(img.height * scale)
    ws.add_image(img, cell)


UI_LEGEND = [
    "1、顶部标题与品牌区", "2、HUD 状态条（SYS/WS/CAM/MODEL/ESP32）",
    "3、侧边导航（控制台 / 数据标注）", "4、摄像头卡片头（设备选择/刷新/预览）",
    "5、实时视频预览区（MJPEG + 检测框叠加）", "6、摄像头运行统计（帧数/客户端/帧率/重连）",
    "7、抓拍训练图区（已抓拍 1974 张）", "8、识别结果卡片头（推理 FPS / 单帧耗时 / 累计帧）",
    "9、多目标策略切换（优先级/最大件/置信度/暂停）", "10、当前识别主结果",
    "11、置信度表盘", "12、类别 → 物理桶明细表",
    "13、分拣统计图（环形=物理桶占比，条形=类别次数）", "14、近 10 分钟分拣趋势",
    "15、下发历史列表", "16、模型管理卡片（模型目录/加载/置信度阈值）",
    "17、串口管理卡片（端口/连接/手动分拣）", "18、主题切换（深色/浅色）",
]


def fill_ui_sheet(wb) -> None:
    ws = wb["UI界面"]
    ws._images = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=30):
        for cell in row:
            if isinstance(cell.value, str) and ("、" in cell.value or cell.value.strip().isdigit()):
                cell.value = None

    add_img(ws, SHOTS / "ui-dark-marked.png", "E5", 1700)
    add_img(ws, SHOTS / "crop-camera.png", "AF5", 620)
    add_img(ws, SHOTS / "crop-result.png", "AF22", 620)
    add_img(ws, SHOTS / "crop-model.png", "AF45", 500)
    add_img(ws, SHOTS / "crop-serial.png", "AF63", 500)
    add_img(ws, SHOTS / "ui-light-marked.png", "AF81", 700)

    ws["B46"] = "编号说明（对应上图中红色圆圈数字）"
    legend_font = copy(ws["B46"].font)
    legend_font.bold = True
    ws["B46"].font = legend_font
    for i, text in enumerate(UI_LEGEND):
        ws[f"B{47 + i}"] = text


RUN_SHOTS = [
    ("图一  控制台运行中（摄像头 Camera 0 已打开 1280×720，模型 trash_final0.pt 已加载）", "ui-dark.png", "D13", 1350),
    ("图二  摄像头卡片细节（设备/刷新列表/停止预览/实时帧统计）", "crop-camera.png", "D45", 1000),
    ("图三  识别结果卡片细节（推理 FPS/单帧耗时/多目标策略/置信度表盘/分桶统计）", "crop-result.png", "D84", 1000),
    ("图四  模型卡片细节（权重列表/加载/置信度阈值/当前模型信息）", "crop-model.png", "P45", 760),
    ("图五  串口卡片细节（端口枚举/连接/手动分拣下发）", "crop-serial.png", "P84", 760),
    ("图六  数据标注页（图片列表 + 画布 + 标注工具）", "ui-annotate.png", "D130", 1350),
]


def fill_run_sheet(wb) -> None:
    ws = wb["运行"]
    ws._images = []
    for caption, fname, cell, width in RUN_SHOTS:
        row = int("".join(ch for ch in cell if ch.isdigit()))
        ws[f"D{row - 2}"] = caption
        add_img(ws, SHOTS / fname, cell, width)


def write_markdown() -> Path:
    """同一份用例的 Markdown 版（便于评审/复制到 Word）。"""
    out = HERE / "测试项目.md"
    rows = CASES_FINAL
    mods: dict[str, list[tuple]] = {}
    for c in rows:
        mods.setdefault(c[0], []).append(c)

    lines = [
        f"# {PROJECT} 项目单体测试书 —— 测试项目",
        "",
        f"- 用例总数：**{len(rows)}** 条（侧重 UI 界面逐项核对 + 实际运行检测）",
        f"- 生成日期：{TODAY}",
        "- 确定结果：`OK` 已实测通过 / `待执行` 需实机补测 / `NG` 不通过",
        "",
        "## 模块分布",
        "",
        "| 模块 | 用例数 |",
        "| --- | --- |",
    ]
    for m, items in mods.items():
        lines.append(f"| {m} | {len(items)} |")
    lines.append("")

    idx = 0
    for m, items in mods.items():
        lines.append(f"## {m}（{len(items)} 条）")
        lines.append("")
        for c in items:
            idx += 1
            _mod, title, pre, steps, exp, method, ev, res, verdict, note = c
            lines.append(f"### {idx}. {title}　`{verdict}`")
            lines.append("")
            lines.append(f"- **前置条件**：{pre}")
            for ln in str(steps).splitlines():
                if ln.strip():
                    lines.append(f"- **操作步骤**：{ln.strip()}")
            lines.append(f"- **预期结果**：{exp}")
            lines.append(f"- **测试方法**：{method}")
            lines.append(f"- **证据**：{ev}")
            if res:
                lines.append(f"- **实测/说明**：{res}")
            if note:
                lines.append(f"- **备注**：{note}")
            lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def main() -> None:
    if not TEMPLATE.exists():
        raise SystemExit(
            "找不到学校模板：\n"
            f"  {TEMPLATE}\n"
            "该模板不入库（已 gitignore）。请把学校下发的模板文件放到项目根目录后重试，\n"
            "文件名必须是：第xx小组_xxx项目_测试成绩书.xlsx"
        )
    shutil.copy2(TEMPLATE, OUT)

    wb = load_workbook(OUT)
    fill_cover(wb)
    n = fill_cases(wb)
    fill_ui_sheet(wb)
    fill_run_sheet(wb)
    wb.save(OUT)
    md = write_markdown()

    ok = sum(1 for c in CASES_FINAL if c[8] == "OK")
    pend = sum(1 for c in CASES_FINAL if c[8] == "待执行")
    ng = sum(1 for c in CASES_FINAL if c[8] == "NG")
    mods: dict[str, int] = {}
    for c in CASES_FINAL:
        mods[c[0]] = mods.get(c[0], 0) + 1

    print(f"已生成：{OUT}")
    print(f"  用例总数 {n}（第 {FIRST_ROW}~{FIRST_ROW + n - 1} 行）")
    print(f"  OK={ok}  待执行={pend}  NG={ng}")
    print(f"  模块分布：{mods}")


if __name__ == "__main__":
    main()
