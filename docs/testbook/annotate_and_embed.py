# -*- coding: utf-8 -*-
"""给 UI 截图自动画编号 + 引线，并插入测试书的「UI界面」「运行」两个表。

输入：docs/testbook/shots/{ui-dark.png, ui-light.png, ui-annotate.png,
        crop-camera.png, crop-result.png, crop-model.png, crop-serial.png, boxes.json}
输出：docs/testbook/shots/ui-dark-marked.png、ui-light-marked.png
      覆盖更新 docs/testbook/垃圾自动识别分类系统_单体测试书.xlsx 的 UI界面 / 运行 表
"""
from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "shots"
XLSX = HERE / "垃圾自动识别分类系统_单体测试书.xlsx"

FONT_NUM = "C:/Windows/Fonts/arialbd.ttf"
BADGE_R = 15
BADGE_FILL = (216, 30, 60)
BADGE_TEXT = (255, 255, 255)
LINE_COLOR = (216, 30, 60)
HALO = (255, 255, 255)


def load_font(size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(FONT_NUM, size)
    except Exception:  # noqa: BLE001
        return ImageFont.load_default()


# 部分编号强制方位，避免引线跨过卡片
SIDE_HINTS = {1: 'below', 2: 'below', 4: 'left', 8: 'right', 16: 'left', 17: 'left', 18: 'below'}


def clamp(v: float, lo: float, hi: float) -> float:
    return max(lo, min(hi, v))


def overlap_area(cx: float, cy: float, rects: list[dict]) -> float:
    """编号圆与外框集合的重叠面积（像素²）。"""
    bx1, by1, bx2, by2 = cx - BADGE_R, cy - BADGE_R, cx + BADGE_R, cy + BADGE_R
    total = 0.0
    for r in rects:
        x1, y1 = r["x"], r["y"]
        x2, y2 = x1 + r["w"], y1 + r["h"]
        w = max(0.0, min(bx2, x2) - max(bx1, x1))
        h = max(0.0, min(by2, y2) - max(by1, y1))
        total += w * h
    return total


def place_badge(box: dict, width: int, height: int,
                obstacles: list[dict], placed: list[dict], side: str | None = None) -> tuple[float, float, float, float]:
    """给编号选一个不压住界面文字的位置，并返回引线锚点。

    候选位置：左、右、上、下（多种间距），按“压住其它控件/已放编号的面积最小”打分。
    这样标题、按钮文字就不会被编号盖住。
    """
    bx, by, bw, bh = box["x"], box["y"], box["w"], box["h"]
    anchor_x = bx + (bw * 0.18 if bw > 400 else bw * 0.5)
    anchor_y = by + bh * 0.5
    cx_mid, cy_mid = bx + bw * 0.5, by + bh * 0.5

    others = [o for o in obstacles if o is not box]
    # 只避让“文字级”小元素（按钮、标签、状态块）；卡片/表头这类大背景不算，
    # 否则编号会被挤到很远、引线又长又乱。
    small = [o for o in others if o["w"] * o["h"] < 45000]
    best = None
    for gap in (24, 34, 48, 70):
        all_sides = {
            'left': (bx - gap, anchor_y),
            'right': (bx + bw + gap, anchor_y),
            'top': (cx_mid, by - gap),
            'below': (cx_mid, by + bh + gap),
            'topleft': (bx - gap, by - gap),
            'bottomright': (bx + bw + gap, by + bh + gap),
        }
        if side:
            candidates = [all_sides[side]]
        else:
            candidates = list(all_sides.values())
        for cx, cy in candidates:
            if not (BADGE_R + 4 <= cx <= width - BADGE_R - 4):
                continue
            if not (BADGE_R + 4 <= cy <= height - BADGE_R - 4):
                continue
            score = overlap_area(cx, cy, small) * 4.0 + overlap_area(cx, cy, placed) * 12.0
            score += math.hypot(cx - anchor_x, cy - anchor_y) * 0.9
            if best is None or score < best[0]:
                best = (score, cx, cy)
    if best is None:
        return clamp(anchor_x, 20, width - 20), clamp(anchor_y, 20, height - 20), anchor_x, anchor_y
    return best[1], best[2], anchor_x, anchor_y


def draw_badge(draw: ImageDraw.ImageDraw, cx: float, cy: float, no: int, font) -> None:
    r = BADGE_R
    draw.ellipse([cx - r - 2, cy - r - 2, cx + r + 2, cy + r + 2], fill=HALO)
    draw.ellipse([cx - r, cy - r, cx + r, cy + r], fill=BADGE_FILL, outline=HALO)
    text = str(no)
    tb = draw.textbbox((0, 0), text, font=font)
    tw, th = tb[2] - tb[0], tb[3] - tb[1]
    draw.text((cx - tw / 2 - tb[0], cy - th / 2 - tb[1]), text, font=font, fill=BADGE_TEXT)


def annotate(src: Path, dst: Path, boxes: list[dict]) -> dict:
    img = Image.open(src).convert("RGB")
    width, height = img.size
    draw = ImageDraw.Draw(img)
    font = load_font(19)

    obstacles = [b["box"] for b in boxes if b.get("box")]
    placed: list[dict] = []
    for item in boxes:
        box = item.get("box")
        if not box or box.get("w", 0) < 6 or box.get("h", 0) < 6:
            continue
        bx, by, ax, ay = place_badge(box, width, height, obstacles, placed, SIDE_HINTS.get(int(item["no"])))
        placed.append({"x": bx - BADGE_R, "y": by - BADGE_R, "w": BADGE_R * 2, "h": BADGE_R * 2})
        # 引线：从编号边缘指向锚点
        dx, dy = ax - bx, ay - by
        dist = math.hypot(dx, dy) or 1
        sx, sy = bx + dx / dist * (BADGE_R + 3), by + dy / dist * (BADGE_R + 3)
        draw.line([sx, sy, ax, ay], fill=HALO, width=4)
        draw.line([sx, sy, ax, ay], fill=LINE_COLOR, width=2)
        draw.ellipse([ax - 3, ay - 3, ax + 3, ay + 3], fill=LINE_COLOR, outline=HALO)
        draw_badge(draw, bx, by, int(item["no"]), font)

    img.save(dst, "PNG", optimize=True)
    return {"file": dst.name, "w": width, "h": height, "marked": sum(1 for b in boxes if b.get("box"))}


def fit(width_px: int, height_px: int, max_w: int) -> tuple[int, int]:
    if width_px <= max_w:
        return width_px, height_px
    scale = max_w / width_px
    return int(width_px * scale), int(height_px * scale)


def add_picture(ws, path: Path, cell: str, max_w: int) -> None:
    img = XLImage(str(path))
    img.width, img.height = fit(img.width, img.height, max_w)
    ws.add_image(img, cell)


def caption(ws, cell: str, text: str) -> None:
    ws[cell] = text
    ws[cell].font = Font(name="微软雅黑", size=10, bold=True, color="C00000")
    ws[cell].alignment = Alignment(horizontal="left", vertical="center")


def main() -> None:
    meta = json.loads((SHOTS / "boxes.json").read_text(encoding="utf-8"))
    boxes = meta["boxes"]

    d = annotate(SHOTS / "ui-dark.png", SHOTS / "ui-dark-marked.png", boxes)
    l = annotate(SHOTS / "ui-light.png", SHOTS / "ui-light-marked.png", boxes)
    print(f"已标注：{d['file']} {d['w']}x{d['h']}（{d['marked']} 个编号）")
    print(f"已标注：{l['file']} {l['w']}x{l['h']}（{l['marked']} 个编号）")

    wb = load_workbook(XLSX)

    # ---------------- UI界面 表 ----------------
    ws = wb["UI界面"]
    ws._images.clear()
    # 找到“截图区”所在的合并单元格行
    start = None
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=1).value
        if isinstance(v, str) and "截图区" in v:
            start = row
            break
    if start is None:
        start = ws.max_row + 2
    ws.cell(row=start, column=1).value = (
        "截图区（下列截图由程序自动截取并标注编号，编号与上表“编号”列一一对应；"
        "如需更高的截图请见 docs/testbook/shots/ 目录）"
    )

    r = start + 1
    caption(ws, f"A{r}", "UI界面图（深色主题 · 控制台全景 · 编号 1~18）")
    add_picture(ws, SHOTS / "ui-dark-marked.png", f"A{r + 1}", 1500)
    r += 1 + math.ceil(1787 * (1500 / 1600) / 20) + 1

    caption(ws, f"A{r}", "UI界面图（浅色主题 · 控制台全景 · 编号 1~18）")
    add_picture(ws, SHOTS / "ui-light-marked.png", f"A{r + 1}", 1500)
    r += 1 + math.ceil(1787 * (1500 / 1600) / 20) + 1

    caption(ws, f"A{r}", "UI界面图（数据标注页）")
    add_picture(ws, SHOTS / "ui-annotate.png", f"A{r + 1}", 1500)

    # ---------------- 运行 表 ----------------
    ws2 = wb["运行"]
    ws2._images.clear()
    start2 = None
    for row in range(1, ws2.max_row + 1):
        v = ws2.cell(row=row, column=1).value
        if isinstance(v, str) and "截图区" in v:
            start2 = row
            break
    if start2 is None:
        start2 = ws2.max_row + 2
    ws2.cell(row=start2, column=1).value = (
        "截图区（实际运行截图；对应“运行图”编号见表内。硬件相关截图需实机补拍）"
    )

    r2 = start2 + 1
    caption(ws2, f"A{r2}", "运行图4 / 运行图5：控制台运行中（摄像头已打开 Camera 0 · 1280×720 · 模型已加载 trash_final0.pt）")
    add_picture(ws2, SHOTS / "ui-dark.png", f"A{r2 + 1}", 1500)
    r2 += 1 + math.ceil(1787 * (1500 / 1600) / 20) + 1

    caption(ws2, f"A{r2}", "运行图：摄像头卡片细节（设备选择 / 刷新列表 / 停止预览 / 实时帧统计）")
    add_picture(ws2, SHOTS / "crop-camera.png", f"A{r2 + 1}", 1200)
    r2 += 1 + math.ceil(747 * (1200 / 768) / 20) + 1

    caption(ws2, f"A{r2}", "运行图：识别结果卡片细节（推理 FPS / 单帧耗时 / 多目标策略 / 置信度表盘 / 分桶统计）")
    add_picture(ws2, SHOTS / "crop-result.png", f"A{r2 + 1}", 1200)
    r2 += 1 + math.ceil(879 * (1200 / 769) / 20) + 1

    caption(ws2, f"A{r2}", "运行图：模型卡片与串口卡片细节")
    add_picture(ws2, SHOTS / "crop-model.png", f"A{r2 + 1}", 900)
    add_picture(ws2, SHOTS / "crop-serial.png", f"J{r2 + 1}", 900)

    # 行高：给图片区域留出足够空间
    for sheet, rows_ in ((ws, (start, r)), (ws2, (start2, r2))):
        for row in range(rows_[0] + 1, rows_[1] + 30):
            sheet.row_dimensions[row].height = 20

    wb.save(XLSX)
    print(f"已写入图片到：{XLSX}")


if __name__ == "__main__":
    main()
